#!/usr/bin/env python3
"""Historical Migration Script: loads 小组奉献资料库_corrected.xlsx and writes transactions to Firestore.

Usage:
    python pipeline/scripts/migrate_xlsx_to_firestore.py [--xlsx PATH] [--project PROJECT] [--dry-run]
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
from datetime import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional
from zoneinfo import ZoneInfo

# Add pipeline directory to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from ledger_pipeline.firestore_db import FirestoreLedgerDB
from ledger_pipeline.ledger import parse_date, parse_money, parse_sheet_name

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("migration")


def find_default_xlsx_path() -> Optional[str]:
    candidates = [
        os.environ.get("XLSX_PATH"),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "小组奉献资料库_corrected.xlsx")),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "小组奉献资料库_corrected.xlsx")),
        os.path.abspath("小组奉献资料库_corrected.xlsx"),
    ]
    for p in candidates:
        if p and os.path.exists(p):
            return p
    return None


def extract_transactions_from_workbook(xlsx_path: str) -> List[Dict[str, Any]]:
    """Reads all month sheets from the workbook and extracts transaction dicts."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    transactions: List[Dict[str, Any]] = []
    tz = ZoneInfo("Asia/Kuala_Lumpur")
    now_iso = datetime.now(tz).isoformat()

    for ws in wb.worksheets:
        ym = parse_sheet_name(ws.title)
        if ym is None:
            continue
        year, month = ym

        for row in range(6, ws.max_row + 1):
            date_val = ws.cell(row, 1).value
            item_val = ws.cell(row, 2).value
            inc_val = ws.cell(row, 3).value
            exp_val = ws.cell(row, 4).value
            bal_val = ws.cell(row, 5).value

            item_str = str(item_val or "").strip()
            try:
                inc = float(parse_money(inc_val) or 0)
            except Exception:
                inc = 0.0
            try:
                exp = float(parse_money(exp_val) or 0)
            except Exception:
                exp = 0.0
            try:
                bal = float(parse_money(bal_val) or 0)
            except Exception:
                bal = 0.0

            # Skip phantom / completely blank rows
            if not item_str and inc == 0.0 and exp == 0.0:
                continue

            parsed_d = parse_date(date_val, year)
            if parsed_d:
                date_str = parsed_d.strftime("%Y-%m-%d")
            elif date_val is not None:
                date_str = str(date_val).strip()
            else:
                date_str = f"{year}-{month:02d}-01"

            transactions.append({
                "date": date_str,
                "item": item_str,
                "income": round(inc, 2),
                "expense": round(exp, 2),
                "balance": round(bal, 2),
                "year": year,
                "month": month,
                "user": "migration",
                "created_at": now_iso,
            })

    # Sort transactions chronologically
    transactions.sort(key=lambda x: (x["date"], x["year"], x["month"]))
    return transactions


def migrate_xlsx_to_firestore(xlsx_path: str, db: Optional[FirestoreLedgerDB] = None,
                              dry_run: bool = False, project_id: Optional[str] = None) -> Dict[str, Any]:
    """Loads transactions from Excel and writes them into Firestore."""
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Workbook not found at {xlsx_path}")

    log.info("Reading transactions from %s ...", xlsx_path)
    transactions = extract_transactions_from_workbook(xlsx_path)
    log.info("Found %d valid transactions across month sheets.", len(transactions))

    if dry_run:
        log.info("Dry-run mode: skipping write to Firestore.")
        return {
            "status": "dry_run",
            "transactions_count": len(transactions),
            "sample": transactions[:5],
        }

    db = db or FirestoreLedgerDB(project_id)
    # Check if real firestore client is connected
    client = getattr(db, "_client", None)
    written = 0

    if client:
        # Batch write in chunks of 450 (Firestore limit is 500 per batch)
        chunk_size = 450
        for i in range(0, len(transactions), chunk_size):
            chunk = transactions[i:i + chunk_size]
            batch = client.batch()
            col = client.collection("transactions")
            for tx in chunk:
                ref = col.document()
                batch.set(ref, tx)
            batch.commit()
            written += len(chunk)
            log.info("Committed batch: %d / %d transactions", written, len(transactions))
    else:
        # In-memory store fallback (tests or offline)
        for tx in transactions:
            db.add_transaction(
                date_str=tx["date"],
                item=tx["item"],
                income=tx["income"],
                expense=tx["expense"],
                balance=tx["balance"],
                user=tx.get("user", "migration"),
            )
            written += 1
        log.info("Wrote %d transactions to in-memory store.", written)

    return {
        "status": "ok",
        "transactions_migrated": written,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Migrate Excel ledger to Cloud Firestore")
    parser.add_argument("--xlsx", help="Path to 小组奉献资料库_corrected.xlsx")
    parser.add_argument("--project", help="Google Cloud project ID", default=os.environ.get("GOOGLE_CLOUD_PROJECT", "group-511516"))
    parser.add_argument("--dry-run", action="store_true", help="Parse and log without writing to Firestore")
    args = parser.parse_args()

    xlsx_path = args.xlsx or find_default_xlsx_path()
    if not xlsx_path:
        log.error("Could not find 小组奉献资料库_corrected.xlsx. Specify with --xlsx.")
        return 1

    try:
        res = migrate_xlsx_to_firestore(xlsx_path, dry_run=args.dry_run, project_id=args.project)
        log.info("Migration finished successfully: %s", res)
        return 0
    except Exception as exc:
        log.exception("Migration failed: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
