"""Where ledger data comes from."""
from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Protocol, Tuple

import requests

from .ledger import RawSheet, parse_sheet_name, sheet_name, ym_next


class Source(Protocol):
    def load(self) -> Dict[str, RawSheet]: ...


class XlsxSource:
    """Local workbook (used for preview/tests and for validating the corrected file)."""

    def __init__(self, path: str):
        self.path = path

    def load(self) -> Dict[str, RawSheet]:
        import openpyxl  # dev/test dependency only
        wb = openpyxl.load_workbook(self.path, data_only=True)
        out: Dict[str, RawSheet] = {}
        for ws in wb.worksheets:
            if parse_sheet_name(ws.title) is None:
                continue
            rows = [[c.value for c in r[:5]] for r in ws.iter_rows(min_row=6, max_col=5)]
            out[ws.title] = RawSheet(
                name=ws.title,
                group_id=str(ws["B1"].value or ""),
                workers=str(ws["B2"].value or ""),
                e5=ws["E5"].value,
                rows=rows,
            )
        return out


class FirestoreSource:
    """Loads ledger data from Cloud Firestore transactions into RawSheets."""

    def __init__(self, db: Optional[Any] = None, opening_balance: Any = 693):
        if db is None:
            from .firestore_db import FirestoreLedgerDB
            self.db = FirestoreLedgerDB()
        else:
            self.db = db
        self.opening_balance = Decimal(str(opening_balance if opening_balance is not None else 693))

    def load(self) -> Dict[str, RawSheet]:
        txs = self.db.query_transactions(limit=None)
        if not txs:
            return {}

        grouped: Dict[Tuple[int, int], List[Dict[str, Any]]] = {}
        for t in txs:
            y = t.get("year")
            m = t.get("month")
            if not y or not m:
                d_str = str(t.get("date", ""))
                parts = d_str.split("-")
                if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
                    y, m = int(parts[0]), int(parts[1])
            if y and m:
                grouped.setdefault((int(y), int(m)), []).append(t)

        if not grouped:
            return {}

        out: Dict[str, RawSheet] = {}
        sorted_yms = sorted(grouped.keys())
        first_ym = sorted_yms[0]
        last_ym = sorted_yms[-1]

        app_settings = self.db.get_app_settings() if hasattr(self.db, "get_app_settings") else {}
        booklet_cfg = (app_settings.get("booklet") if isinstance(app_settings, dict) else None) or {}
        group_id = booklet_cfg.get("group_code") or "YA5"
        workers = booklet_cfg.get("counters") or "Rivaldo, Brendon"

        running_bal = self.opening_balance
        curr_ym = first_ym
        while curr_ym <= last_ym:
            y, m = curr_ym
            s_name = sheet_name(y, m)
            month_txs = grouped.get(curr_ym, [])
            e5_val = running_bal

            # Sort transactions in this month chronologically
            month_txs.sort(key=lambda x: (str(x.get("date", "")), str(x.get("created_at", ""))))

            rows: List[List[Any]] = []
            for tx in month_txs:
                inc = Decimal(str(tx.get("income", 0) or 0))
                exp = Decimal(str(tx.get("expense", 0) or 0))
                running_bal = running_bal + inc - exp
                rows.append([
                    tx.get("date"),
                    tx.get("item", ""),
                    float(inc) if inc > 0 else None,
                    float(exp) if exp > 0 else None,
                    float(running_bal),
                ])

            if month_txs:
                out[s_name] = RawSheet(
                    name=s_name,
                    group_id=group_id,
                    workers=workers,
                    e5=float(e5_val),
                    rows=rows,
                )

            curr_ym = ym_next(curr_ym)

        return out

