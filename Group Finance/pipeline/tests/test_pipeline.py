import base64
import hashlib
import hmac
import json
import os
import unittest
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

HERE = os.path.dirname(__file__)
import sys
sys.path.insert(0, os.path.join(HERE, ".."))

from ledger_pipeline.config import Settings
from ledger_pipeline.ledger import (RawSheet, build_months, parse_date, parse_money, parse_sheet_name)
from ledger_pipeline.pipeline import Pipeline, parse_ym
from ledger_pipeline.render import render_all_months_html, render_month_card, render_month_html
from ledger_pipeline.security import verify_netlify_signature, verify_token
from ledger_pipeline.storage import LocalStorage
import tempfile

CORRECTED = os.path.join(HERE, "..", "..", "小组奉献资料库_corrected.xlsx")


def sheets_fixture():
    return {
        "Aug 2026": RawSheet("Aug 2026", "YA5", "Rivaldo, Brendon", 498,
                             [["2026-08-13", "小组", 90, 0, 588]]),
        "Sep 2026": RawSheet("Sep 2026", "YA5", "Rivaldo, Brendon", datetime(2026, 8, 30),   # garbage E5
                             [["2026-09-09", "小组", 120, 0, None],
                              ["2026-09-07", "小组茶点", 0, 50, None],        # out-of-order input
                              [None, None, None, None, None]]),
    }


class LedgerTests(unittest.TestCase):
    def test_chain_and_sorting(self):
        ms = build_months(sheets_fixture(), upto=(2026, 10))
        aug, sep, octo = ms[(2026, 8)], ms[(2026, 9)], ms[(2026, 10)]
        self.assertEqual(aug.opening, 498)
        self.assertEqual(aug.closing, 588)
        self.assertEqual(sep.opening, 588)                     # taken from Aug, garbage E5 ignored
        self.assertEqual([e.item for e in sep.entries], ["小组茶点", "小组"])   # chronological
        self.assertEqual([e.balance for e in sep.entries], [538, 658])
        self.assertEqual(sep.closing, 658)
        self.assertFalse(octo.sheet_exists)                    # month without sheet yet
        self.assertEqual(octo.opening, 658)                    # b/f carried
        self.assertEqual(octo.entries, [])
        self.assertTrue(any("E5 is not a number" in w for w in sep.warnings))

    def test_first_month_uses_configured_opening(self):
        ms = build_months({"Jan 2026": RawSheet("Jan 2026", rows=[["2026-01-14", "x", 10, 0, None]])},
                          opening_balance=Decimal(693))
        self.assertEqual(ms[(2026, 1)].closing, 703)

    def test_decimals_and_year_rollover(self):
        s = {"Dec 2026": RawSheet("Dec 2026", e5=10, rows=[["2026-12-31", "a", "0.10", 0, None],
                                                           ["2026-12-31", "b", 0, "0.30", None]])}
        ms = build_months(s, upto=(2027, 1))
        self.assertEqual(ms[(2026, 12)].closing, Decimal("9.80"))
        self.assertEqual(ms[(2027, 1)].opening, Decimal("9.80"))

    def test_parsers(self):
        self.assertEqual(parse_date(46272), date(2026, 9, 7))       # sheets serial
        self.assertEqual(parse_date(46247), date(2026, 8, 13))
        self.assertEqual(parse_date("13/08", 2026), date(2026, 8, 13))
        self.assertEqual(parse_date("2026-08-13T00:00:00.000Z"), date(2026, 8, 13))
        self.assertEqual(parse_money("RM 1,200.50"), Decimal("1200.50"))
        self.assertIsNone(parse_money(""))
        with self.assertRaises(ValueError):
            parse_money(datetime(2026, 1, 1))
        self.assertEqual(parse_sheet_name("Sep 2026"), (2026, 9))
        self.assertIsNone(parse_sheet_name("Template"))

    def test_garbage_amounts_do_not_crash(self):
        s = {"Sep 2026": RawSheet("Sep 2026", e5=0, rows=[["2026-09-01", "x", "abc", 5, None]])}
        m = build_months(s)[(2026, 9)]
        self.assertEqual(m.closing, -5)
        self.assertTrue(any("not a number" in w for w in m.warnings))

    def test_build_entries_ignores_phantom_rows(self):
        s = {"Oct 2026": RawSheet("Oct 2026", e5=319, rows=[
            ["2026-10-14", "", 0, 0, 319],
            ["2026-10-14", "   ", 0, 0, 319],
            ["2026-10-14", "实际奉献", 100, 0, 419],
        ])}
        m = build_months(s)[(2026, 10)]
        self.assertEqual(len(m.entries), 1)
        self.assertEqual(m.entries[0].item, "实际奉献")
        self.assertEqual(m.entries[0].balance, 419)
        self.assertEqual(m.closing, 419)

    @unittest.skipUnless(os.path.exists(CORRECTED), "corrected workbook not generated")
    def test_corrected_workbook_matches_booklet(self):
        from ledger_pipeline.sources import XlsxSource
        ms = build_months(XlsxSource(CORRECTED).load(), upto=(2026, 9))
        want = {1: 512, 2: 452, 3: 598, 4: 610, 5: 660, 6: 710, 7: 498, 8: 588, 9: 319}
        self.assertEqual({m: ms[(2026, m)].closing for m in want}, want)
        for k in want:
            self.assertEqual(ms[(2026, k)].warnings, [], ((2026, k), ms[(2026, k)].warnings))


class RenderTests(unittest.TestCase):
    def test_form_content(self):
        m = build_months(sheets_fixture())[(2026, 8)]
        html = render_month_html(m, closed=True, show_nav=False)
        for needle in ["小组每月奉献记录", "小组编号：", "__YA5___", "点算同工：", "Rivaldo, Brendon", "月份：",
                       "_8 /2026", "数目-收入", "数目-支出", "结存", "签名", "Balance b/f",
                       "08-13", ">588<", "Verify by: ________Brendon__________",
                       "Name     :  _______Rivaldo___________", "Date:  _______31/8/2026_______"]:
            self.assertIn(needle, html)
        self.assertEqual(html.count("<tr>"), 1 + 4)       # header + b/f + 1 entry + 2 blanks
        self.assertNotIn('class="nav"', html)

    def test_open_month_has_blank_signature_block(self):
        m = build_months(sheets_fixture(), upto=(2026, 9))[(2026, 9)]
        html = render_month_html(m, closed=False)
        self.assertNotIn("Brendon__", html)
        self.assertIn("Date:  ________________________", html)

    def test_html_escaped(self):
        s = {"Sep 2026": RawSheet("Sep 2026", e5=0, rows=[["2026-09-01", "<script>x</script>", 1, 0, None]])}
        html = render_month_html(build_months(s)[(2026, 9)])
        self.assertNotIn("<script>x", html)

    def test_render_month_card_editable_and_non_editable(self):
        m = build_months(sheets_fixture())[(2026, 8)]
        card_editable = render_month_card(m, editable=True)
        self.assertIn('<th class="col-act"></th>', card_editable)
        self.assertNotIn("操作", card_editable)
        self.assertIn('class="btn-del"', card_editable)
        self.assertIn('deleteEntry(\'2026\', \'8\', \'08-13\', \'小组\', \'90\', \'0\', this)', card_editable)

        card_readonly = render_month_card(m, editable=False)
        self.assertNotIn('col-act', card_readonly)
        self.assertNotIn('btn-del', card_readonly)

        html_print = render_month_html(m, show_nav=False)
        self.assertNotIn('<th class="col-act">', html_print)
        self.assertNotIn('<button', html_print)
        self.assertNotIn('deleteEntry', html_print)

        html_interactive = render_month_html(m, show_nav=True)
        self.assertIn('<th class="col-act"></th>', html_interactive)
        self.assertNotIn("操作", html_interactive)
        self.assertIn('<button type="button" class="btn-del"', html_interactive)
        self.assertIn('async function deleteEntry', html_interactive)

        all_html = render_all_months_html({(2026, 8): m}, current_ym=(2026, 8), closed_map={(2026, 8): False})
        self.assertIn('col-act', all_html)
        self.assertIn('async function deleteEntry', all_html)



class FakeSource:
    def __init__(self, sheets): self.sheets = sheets
    def load(self): return self.sheets


class FakeDrive:
    def __init__(self): self.saved = []
    def save_pdf(self, fn, pdf):
        self.saved.append((fn, pdf)); return {"id": "abc", "url": "https://drive/x"}


def at(y, m, d, h=12):
    return lambda: datetime(y, m, d, h, 0, tzinfo=ZoneInfo("Asia/Kuala_Lumpur"))


class PipelineTests(unittest.TestCase):
    def make(self, now, sheets=None):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.drive = FakeDrive()
        s = Settings(pipeline_token="tok", page_cache_seconds=60)
        return Pipeline(s, FakeSource(sheets or sheets_fixture()), LocalStorage(self.tmp.name),
                        drive=self.drive, pdf=lambda h: b"%PDF-" + h.encode(), now=now)

    def read(self, name):
        with open(os.path.join(self.tmp.name, name), encoding="utf-8") as f:
            return f.read()

    def test_sync_creates_current_month_even_without_sheet(self):
        p = self.make(at(2026, 10, 14))
        r = p.publish_site()
        self.assertEqual(r["current"], "2026-10")
        self.assertEqual(r["opening"], 658)                      # last month's balance filled in b/f
        idx = self.read("index.html")
        self.assertIn("_10 /2026", idx)
        self.assertIn(">658<", idx)
        self.assertNotIn("Brendon__", self.read("2026-10.html")) # open month not signed
        self.assertIn("Brendon__", self.read("2026-09.html"))    # past month signed
        self.assertIn("31/8/2026", self.read("2026-08.html"))
        self.assertIn("30/9/2026", self.read("2026-09.html"))
        self.assertEqual(json.loads(self.read("months.json"))["current"], "2026-10")

    def test_close_skips_when_not_last_day(self):
        p = self.make(at(2026, 9, 29))
        self.assertEqual(p.close_month()["status"], "skipped")
        self.assertEqual(self.drive.saved, [])

    def test_close_on_last_day(self):
        p = self.make(at(2026, 9, 30, 23))
        r = p.close_month()
        self.assertEqual(r["status"], "ok")
        self.assertEqual(r["closing_date"], "2026-09-30")
        self.assertEqual(r["closing"], 658)
        fn, pdf = self.drive.saved[0]
        self.assertEqual(fn, "小组奉献记录_2026-09.pdf")
        self.assertIn("Date:  _______30/9/2026_______", pdf.decode())
        self.assertIn("Name     :  _______Rivaldo___________", pdf.decode())
        # same day: marker makes the live page show the signed block
        self.assertIn("30/9/2026", self.read("2026-09.html"))
        self.assertTrue(os.path.exists(os.path.join(self.tmp.name, "archive/小组奉献记录_2026-09.pdf")))

    def test_close_explicit_past_month(self):
        p = self.make(at(2026, 10, 14))
        r = p.close_month(parse_ym("2026-08"))
        self.assertEqual(r["closing"], 588)
        self.assertIn("31/8/2026", self.drive.saved[0][1].decode())

    def test_archive_and_state_use_private_storage(self):
        site, priv = tempfile.TemporaryDirectory(), tempfile.TemporaryDirectory()
        self.addCleanup(site.cleanup); self.addCleanup(priv.cleanup)
        p = Pipeline(Settings(), FakeSource(sheets_fixture()), LocalStorage(site.name), drive=FakeDrive(),
                     pdf=lambda h: b"%PDF-" + h.encode(), now=at(2026, 9, 30, 23), private=LocalStorage(priv.name))
        p.close_month()
        self.assertTrue(os.path.exists(os.path.join(priv.name, "archive/小组奉献记录_2026-09.pdf")))
        self.assertTrue(os.path.exists(os.path.join(priv.name, "state/closed-2026-09.json")))
        self.assertFalse(os.path.exists(os.path.join(site.name, "archive")))
        self.assertFalse(os.path.exists(os.path.join(site.name, "state")))
        self.assertTrue(os.path.exists(os.path.join(site.name, "2026-09.html")))

    def test_year_end_close(self):
        s = {"Dec 2026": RawSheet("Dec 2026", e5=5, rows=[["2026-12-31", "x", 1, 0, 6]])}
        p = self.make(at(2026, 12, 31, 23), s)
        self.assertEqual(p.close_month()["closing"], 6)
        p2 = self.make(at(2027, 1, 1, 1), s)
        self.assertEqual(p2.publish_site()["opening"], 6)


def jws(secret, body, iss="netlify"):
    b = lambda d: base64.urlsafe_b64encode(d).rstrip(b"=").decode()
    h = b(json.dumps({"alg": "HS256", "typ": "JWT"}).encode())
    p = b(json.dumps({"iss": iss, "sha256": hashlib.sha256(body).hexdigest()}).encode())
    sig = b(hmac.new(secret.encode(), f"{h}.{p}".encode(), hashlib.sha256).digest())
    return f"{h}.{p}.{sig}"


class SecurityAndServiceTests(unittest.TestCase):
    def test_netlify_signature(self):
        body = b'{"state":"ready"}'
        self.assertTrue(verify_netlify_signature(jws("s3", body), body, "s3"))
        self.assertFalse(verify_netlify_signature(jws("s3", body), body + b"x", "s3"))   # body tampered
        self.assertFalse(verify_netlify_signature(jws("bad", body), body, "s3"))
        self.assertFalse(verify_netlify_signature(jws("s3", body, iss="evil"), body, "s3"))
        self.assertFalse(verify_netlify_signature(jws("", body), body, ""))               # fail closed
        self.assertFalse(verify_netlify_signature("garbage", body, "s3"))
        self.assertFalse(verify_netlify_signature(None, body, "s3"))

    def test_token(self):
        self.assertTrue(verify_token({"X-Pipeline-Token": "t"}, "t"))
        self.assertTrue(verify_token({"Authorization": "Bearer t"}, "t"))
        self.assertFalse(verify_token({"X-Pipeline-Token": "x"}, "t"))
        self.assertFalse(verify_token({}, "t"))
        self.assertFalse(verify_token({"X-Pipeline-Token": ""}, ""))

    def test_endpoints(self):
        from ledger_pipeline.service import create_app
        tmp = tempfile.TemporaryDirectory(); self.addCleanup(tmp.cleanup)
        drive = FakeDrive()
        s = Settings(pipeline_token="tok", netlify_jws_secret="nf")
        p = Pipeline(s, FakeSource(sheets_fixture()), LocalStorage(tmp.name), drive=drive,
                     pdf=lambda h: b"pdf", now=at(2026, 10, 14))
        c = create_app(s, p).test_client()
        self.assertEqual(c.get("/healthz").status_code, 200)
        self.assertEqual(c.post("/sync").status_code, 401)
        self.assertEqual(c.post("/sync", headers={"X-Pipeline-Token": "wrong"}).status_code, 401)
        self.assertEqual(c.post("/sync", headers={"X-Pipeline-Token": "tok"}).status_code, 200)
        body = b'{"id":"1"}'
        r = c.post("/sync", data=body, headers={"X-Webhook-Signature": jws("nf", body)})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["current"], "2026-10")
        # netlify signature must NOT authorise /close
        self.assertEqual(c.post("/close", data=body, headers={"X-Webhook-Signature": jws("nf", body)}).status_code, 401)
        r = c.post("/close", headers={"X-Pipeline-Token": "tok"})        # 14 Oct is not the last day
        self.assertEqual(r.get_json()["status"], "skipped")
        r = c.post("/close?month=2026-09", headers={"X-Pipeline-Token": "tok"})
        self.assertEqual(r.get_json()["status"], "ok")
        self.assertEqual(len(drive.saved), 1)
        self.assertEqual(c.post("/close?month=2026-13", headers={"X-Pipeline-Token": "tok"}).status_code, 400)

        # /api tests
        # 1. Login
        r_login_fail = c.post("/api", json={"action": "login", "username": "admin", "password": "wrongpassword"})
        self.assertEqual(r_login_fail.status_code, 401)
        r_login_ok = c.post("/api", json={"action": "login", "username": "admin", "password": "Abc1234#"})
        self.assertEqual(r_login_ok.status_code, 200)
        self.assertEqual(r_login_ok.get_json()["status"], "success")

        # 2. Yearly summary
        r_sum = c.post("/api", json={"action": "get_yearly_summary", "year": 2026})
        self.assertEqual(r_sum.status_code, 200)
        sum_data = r_sum.get_json()
        self.assertEqual(sum_data["status"], "success")
        self.assertEqual(len(sum_data["data"]), 12)
        self.assertIn("totalIncome", sum_data)
        self.assertIn("balance", sum_data)

        # 3. Add record
        r_add = c.post("/api", json={"date": "2026-10-14", "item": "测试同工会", "income": 50, "expense": 0})
        self.assertEqual(r_add.status_code, 200)
        add_data = r_add.get_json()
        self.assertEqual(add_data["status"], "success")
        # 4. Delete record
        r_del = c.post("/api", json={"action": "delete_record", "year": 2026, "month": 10,
                                     "date": "2026-10-14", "item": "测试同工会", "income": 50, "expense": 0})
        self.assertEqual(r_del.status_code, 200)
        self.assertEqual(r_del.get_json()["status"], "success")

        # 5. Route check
        self.assertEqual(c.get("/manage").status_code, 200)
        self.assertEqual(c.get("/").status_code, 200)

    def test_get_xlsx_path(self):
        import ledger_pipeline.service as srv
        from unittest.mock import patch

        # 1. Pipeline source path
        class DummySource:
            path = "/explicit/source.xlsx"
        class DummyPipeline:
            source = DummySource()
        self.assertEqual(srv.get_xlsx_path(DummyPipeline()), "/explicit/source.xlsx")

        # 2. env_path when exists
        with patch.dict(os.environ, {"XLSX_PATH": "/custom/env.xlsx"}):
            with patch("os.path.exists", side_effect=lambda p: p == "/custom/env.xlsx"):
                self.assertEqual(srv.get_xlsx_path(), "/custom/env.xlsx")

        # 3. pipe_path when env_path does not exist
        pipe_expected = os.path.abspath(os.path.join(os.path.dirname(srv.__file__), "..", "小组奉献资料库_corrected.xlsx"))
        with patch.dict(os.environ, {}, clear=True):
            with patch("os.path.exists", side_effect=lambda p: p == pipe_expected):
                self.assertEqual(srv.get_xlsx_path(), pipe_expected)

        # 4. default_path when pipe_path does not exist
        default_expected = os.path.abspath(os.path.join(os.path.dirname(srv.__file__), "..", "..", "小组奉献资料库_corrected.xlsx"))
        with patch.dict(os.environ, {}, clear=True):
            with patch("os.path.exists", side_effect=lambda p: p == default_expected):
                self.assertEqual(srv.get_xlsx_path(), default_expected)


class RecordDeletionTests(unittest.TestCase):
    def test_firestore_delete_transaction(self):
        from ledger_pipeline.firestore_db import FirestoreLedgerDB
        db = FirestoreLedgerDB("dummy-project")
        t1 = db.add_transaction(date_str="2026-08-13", item="项目A", income=100.0, expense=0.0, balance=500.0)
        t2 = db.add_transaction(date_str="2026-08-15", item="项目B", income=0.0, expense=50.0, balance=450.0)
        self.assertEqual(len(db.get_recent_transactions()), 2)

        # Delete non-existent
        deleted = db.delete_transaction(year=2026, month=8, date_str="2026-08-13", item="不存在", income=100.0, expense=0.0)
        self.assertFalse(deleted)
        self.assertEqual(len(db.get_recent_transactions()), 2)

        # Delete t1
        deleted = db.delete_transaction(year=2026, month=8, date_str="2026-08-13", item="项目A", income=100.0, expense=0.0)
        self.assertTrue(deleted)
        recent = db.get_recent_transactions()
        self.assertEqual(len(recent), 1)
        self.assertEqual(recent[0][1], "项目B")

    def test_xlsx_source_delete_row(self):
        import openpyxl
        from ledger_pipeline.service import create_app
        from ledger_pipeline.sources import XlsxSource
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        xlsx_path = os.path.join(tmp.name, "test.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Aug 2026"
        ws["B1"] = "YA5"
        ws["B2"] = "Rivaldo, Brendon"
        ws["E5"] = 498
        ws.append([datetime(2026, 8, 13), "小组奉献", 90, 0, None])
        ws.append([datetime(2026, 8, 20), "活动支出", 0, 30, None])
        wb.save(xlsx_path)

        s = Settings(pipeline_token="tok")
        src = XlsxSource(xlsx_path)
        p = Pipeline(s, src, LocalStorage(tmp.name), pdf=lambda h: b"pdf", now=at(2026, 8, 25))
        c = create_app(s, p).test_client()

        # Delete "小组奉献"
        r = c.post("/api", json={
            "action": "delete_record",
            "year": 2026,
            "month": 8,
            "date": "08-13",
            "item": "小组奉献",
            "income": 90,
            "expense": 0
        })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["status"], "success")

        # Verify in xlsx that row was removed
        wb2 = openpyxl.load_workbook(xlsx_path)
        ws2 = wb2["Aug 2026"]
        rows = list(ws2.iter_rows(min_row=6, max_col=5, values_only=True))
        valid_rows = [row for row in rows if any(row)]
        self.assertEqual(len(valid_rows), 1)
        self.assertEqual(valid_rows[0][1], "活动支出")
        self.assertEqual(valid_rows[0][3], 30)

    def test_add_record_rejects_empty_item_and_zero_amounts(self):
        from ledger_pipeline.service import create_app
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        s = Settings(pipeline_token="tok")
        p = Pipeline(s, FakeSource(sheets_fixture()), LocalStorage(tmp.name), pdf=lambda h: b"pdf", now=at(2026, 10, 14))
        c = create_app(s, p).test_client()

        # Empty item and 0 amounts
        r = c.post("/api", json={"date": "2026-10-14", "item": "", "income": 0, "expense": 0})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.get_json()["message"], "事项或金额不能为空")

        # Empty item with spaces and 0 amounts
        r2 = c.post("/api", json={"date": "2026-10-14", "item": "   ", "income": 0, "expense": 0})
        self.assertEqual(r2.status_code, 400)
        self.assertEqual(r2.get_json()["message"], "事项或金额不能为空")

    def test_delete_record_from_xlsx_regardless_of_source_type(self):
        import openpyxl
        from ledger_pipeline.service import create_app
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        xlsx_path = os.path.join(tmp.name, "test.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Aug 2026"
        ws["B1"] = "YA5"
        ws["B2"] = "Rivaldo, Brendon"
        ws["E5"] = 498
        ws.append([datetime(2026, 8, 13), "小组奉献", 90, 0, None])
        wb.save(xlsx_path)

        os.environ["XLSX_PATH"] = xlsx_path
        try:
            s = Settings(pipeline_token="tok")
            # p.source is FakeSource, NOT XlsxSource
            p = Pipeline(s, FakeSource(sheets_fixture()), LocalStorage(tmp.name), pdf=lambda h: b"pdf", now=at(2026, 8, 25))
            c = create_app(s, p).test_client()

            r = c.post("/api", json={
                "action": "delete_record",
                "year": 2026,
                "month": 8,
                "date": "08-13",
                "item": "小组奉献",
                "income": 90,
                "expense": 0
            })
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.get_json()["status"], "success")

            wb2 = openpyxl.load_workbook(xlsx_path)
            ws2 = wb2["Aug 2026"]
            rows = [row for row in list(ws2.iter_rows(min_row=6, max_col=5, values_only=True)) if any(row)]
            self.assertEqual(len(rows), 0)
        finally:
            os.environ.pop("XLSX_PATH", None)

    def test_delete_phantom_rows_cleanup_in_xlsx(self):
        import openpyxl
        from ledger_pipeline.service import create_app
        from ledger_pipeline.sources import XlsxSource
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        xlsx_path = os.path.join(tmp.name, "test_phantom.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Oct 2026"
        ws["B1"] = "YA5"
        ws["B2"] = "Rivaldo, Brendon"
        ws["E5"] = 319
        ws.append([datetime(2026, 10, 14), "", 0, 0, 319])
        ws.append([datetime(2026, 10, 14), None, None, None, 319])
        ws.append([datetime(2026, 10, 14), "正常奉献", 100, 0, 419])
        wb.save(xlsx_path)

        s = Settings(pipeline_token="tok")
        src = XlsxSource(xlsx_path)
        p = Pipeline(s, src, LocalStorage(tmp.name), pdf=lambda h: b"pdf", now=at(2026, 10, 20))
        c = create_app(s, p).test_client()

        # Delete phantom row
        r = c.post("/api", json={
            "action": "delete_record",
            "year": 2026,
            "month": 10,
            "date": "10-14",
            "item": "",
            "income": 0,
            "expense": 0
        })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["status"], "success")

        wb2 = openpyxl.load_workbook(xlsx_path)
        ws2 = wb2["Oct 2026"]
        rows = [row for row in list(ws2.iter_rows(min_row=6, max_col=5, values_only=True)) if any(row)]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], "正常奉献")
        self.assertEqual(rows[0][2], 100)

    def test_publish_site_fallback_to_xlsx_on_source_error(self):
        import openpyxl
        from ledger_pipeline.pipeline import Pipeline
        from ledger_pipeline.storage import LocalStorage
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        xlsx_path = os.path.join(tmp.name, "fallback.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Oct 2026"
        ws["B1"] = "YA5"
        ws["B2"] = "Rivaldo, Brendon"
        ws["E5"] = 319
        ws.append([datetime(2026, 10, 14), "回退测试奉献", 50, 0, 369])
        wb.save(xlsx_path)

        class BrokenSource:
            def __init__(self, path):
                self.path = path
            def load(self):
                raise RuntimeError("Apps Script unavailable!")

        s = Settings(pipeline_token="tok")
        p = Pipeline(s, BrokenSource(xlsx_path), LocalStorage(tmp.name), now=at(2026, 10, 14))

        # publish_site should not raise, but fall back to xlsx_path!
        res = p.publish_site()
        self.assertEqual(res["status"], "ok")
        self.assertTrue(os.path.exists(os.path.join(tmp.name, "index.html")))
        with open(os.path.join(tmp.name, "index.html"), "r", encoding="utf-8") as f:
            content = f.read()
        self.assertIn("回退测试奉献", content)

    def test_favicon_and_api_content_type(self):
        from ledger_pipeline.service import create_app
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        s = Settings(pipeline_token="tok")
        p = Pipeline(s, FakeSource(sheets_fixture()), LocalStorage(tmp.name), pdf=lambda h: b"pdf", now=at(2026, 10, 14))
        c = create_app(s, p).test_client()

        # 1. Favicon returns 204
        r_fav = c.get("/favicon.ico")
        self.assertEqual(r_fav.status_code, 204)

        # 2. POST /api with empty body returns 400 with "请求体不能为空。"
        r_empty = c.post("/api")
        self.assertEqual(r_empty.status_code, 400)
        self.assertEqual(r_empty.get_json()["message"], "请求体不能为空。")

        # 3. POST /api without application/json (e.g. text/plain)
        r_login = c.post(
            "/api",
            data=json.dumps({"action": "login", "username": "admin", "password": "Abc1234#"}),
            content_type="text/plain"
        )
        self.assertEqual(r_login.status_code, 200)
        self.assertEqual(r_login.get_json()["status"], "success")

        r_sum = c.post(
            "/api",
            data=json.dumps({"action": "get_yearly_summary", "year": 2026}),
            content_type="text/plain"
        )
        self.assertEqual(r_sum.status_code, 200)
        self.assertEqual(r_sum.get_json()["status"], "success")


if __name__ == "__main__":
    unittest.main()
