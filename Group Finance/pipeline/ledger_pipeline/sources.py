"""Where ledger data comes from."""
from __future__ import annotations

import re
from datetime import date, datetime
from typing import Dict, Optional, Protocol

import requests

from .ledger import RawSheet, parse_sheet_name


class Source(Protocol):
    def load(self) -> Dict[str, RawSheet]: ...


class XlsxSource:
    """Local workbook (used for preview/tests and for validating the corrected file)."""

    def __init__(self, path: str):
        self.path = path

    def load(self) -> Dict[str, RawSheet]:
        import openpyxl  # dev/test dependency only
        wb = openpyxl.load_workbook(self.path, data_only=True)
        out: Dict[str, RawSheet] = {}
        for ws in wb.worksheets:
            if parse_sheet_name(ws.title) is None:
                continue
            rows = [[c.value for c in r[:5]] for r in ws.iter_rows(min_row=6, max_col=5)]
            out[ws.title] = RawSheet(
                name=ws.title,
                group_id=str(ws["B1"].value or ""),
                workers=str(ws["B2"].value or ""),
                e5=ws["E5"].value,
                rows=rows,
            )
        return out


class AppsScriptClient:
    """Talks to the (token-protected) actions of the Apps Script web app."""

    def __init__(self, url: str, token: str, timeout: int = 60, session=None):
        if not url:
            raise ValueError("APPS_SCRIPT_URL is not configured")
        self.url, self.token, self.timeout = url, token, timeout
        self.http = session or requests

    def call(self, action: str, **payload) -> dict:
        # Apps Script answers POSTs with a 302 to script.googleusercontent.com; requests follows it.
        r = self.http.post(self.url, json={"action": action, "token": self.token, **payload},
                           timeout=self.timeout)
        r.raise_for_status()
        try:
            data = r.json()
        except ValueError:
            raise RuntimeError(f"Apps Script returned non-JSON: {r.text[:200]!r}")
        if data.get("status") != "success":
            raise RuntimeError(f"Apps Script error for {action}: {data.get('message')}")
        return data


class AppsScriptSource:
    def __init__(self, client: AppsScriptClient):
        self.client = client

    def load(self) -> Dict[str, RawSheet]:
        data = self.client.call("export_ledger")
        out: Dict[str, RawSheet] = {}
        for name, s in data["sheets"].items():
            out[name] = RawSheet(name=name, group_id=s.get("group_id", ""),
                                 workers=s.get("workers", ""), e5=s.get("e5"),
                                 rows=s.get("rows", []))
        return out
